
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook

from src.documents import ingest_documents as ingest


class TestDocumentIngestion(unittest.TestCase):

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.input_dir = self.root / "data" / "document"
        self.input_dir.mkdir(parents=True)
        self.output = self.root / "reports" / "document_analysis.json"

    def tearDown(self):
        self.temp.cleanup()

    def create_docx(self, name, text):
        path = self.input_dir / name
        doc = Document()
        doc.add_paragraph(text)
        doc.save(path)
        return path

    def create_xlsx(self, name, rows):
        path = self.input_dir / name
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data"

        for row in rows:
            sheet.append(row)

        workbook.save(path)
        return path

    def run_ingestion(self):
        with (
            patch.object(ingest, "ROOT", self.root),
            patch.object(ingest, "INPUT_DIR", self.input_dir),
            patch.object(ingest, "OUTPUT_FILE", self.output),
        ):
            return ingest.main()

    def test_reads_three_docx_and_one_xlsx(self):
        originals = [
            self.create_docx("a.docx", "Proposal A"),
            self.create_docx("b.docx", "Proposal B"),
            self.create_docx("c.docx", "Proposal C"),
            self.create_xlsx("data.xlsx", [["Budget", 100]]),
        ]
        hashes = {p.name: ingest.sha256_file(p) for p in originals}

        self.assertEqual(self.run_ingestion(), 0)
        result = json.loads(self.output.read_text(encoding="utf-8"))

        self.assertEqual(result["document_count"], 4)
        self.assertEqual(result["input_mode"], "real")
        self.assertEqual(len(result["source_files"]), 4)
        self.assertTrue(result["evidence_blocks"])

        actual_hashes = {
            item["file_name"]: item["sha256"]
            for item in result["source_files"]
        }
        self.assertEqual(actual_hashes, hashes)

        self.assertEqual(
            {p.name for p in self.input_dir.iterdir()},
            {p.name for p in originals},
        )

    def test_rejects_wrong_file_count_without_output(self):
        self.create_docx("sample.docx", "Test document")

        self.assertEqual(self.run_ingestion(), 1)
        self.assertFalse(self.output.exists())

    def test_rejects_wrong_document_types(self):
        for name in ("a.docx", "b.docx", "c.docx", "d.docx"):
            self.create_docx(name, name)

        self.assertEqual(self.run_ingestion(), 1)
        self.assertFalse(self.output.exists())

    def test_rejects_empty_document(self):
        self.create_docx("a.docx", "")
        self.create_docx("b.docx", "B")
        self.create_docx("c.docx", "C")
        self.create_xlsx("data.xlsx", [["Budget", 100]])

        self.assertEqual(self.run_ingestion(), 1)
        self.assertFalse(self.output.exists())


if __name__ == "__main__":
    unittest.main()
