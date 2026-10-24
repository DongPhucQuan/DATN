from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document


ROOT = Path(__file__).resolve().parents[2]
DOCUMENTS_DIR = ROOT / "data" / "documents"
REPORTS_DIR = ROOT / "reports"

DEFAULT_MAX_DOCUMENTS = 50

PLACEHOLDER_PATTERNS = [
    re.compile(r"\bTODO\b", re.IGNORECASE),
    re.compile(r"\bTBD\b", re.IGNORECASE),
    re.compile(r"\bFIXME\b", re.IGNORECASE),
    re.compile(r"\bXXX\b"),
    re.compile(r"\[.*?(?:điền|bổ sung|cập nhật).*?\]", re.IGNORECASE),
    re.compile(r"<.*?(?:điền|bổ sung|cập nhật).*?>", re.IGNORECASE),
    re.compile(r"\bchưa\s+(?:xác\s+định|cập\s+nhật|bổ\s+sung)\b", re.IGNORECASE),
]

TOPIC_MAP = {
    "purpose_and_scope": [
        "mục tiêu",
        "mục đích",
        "phạm vi",
        "đối tượng áp dụng",
        "scope",
        "objective",
        "purpose",
    ],
    "requirements": [
        "yêu cầu",
        "chức năng",
        "phi chức năng",
        "requirement",
        "functional",
        "non-functional",
    ],
    "stakeholders": [
        "đơn vị",
        "phòng ban",
        "người dùng",
        "bên liên quan",
        "stakeholder",
        "user",
    ],
    "schedule": [
        "tiến độ",
        "thời gian",
        "giai đoạn",
        "kế hoạch triển khai",
        "deadline",
        "milestone",
        "schedule",
    ],
    "budget": [
        "kinh phí",
        "ngân sách",
        "chi phí",
        "dự toán",
        "budget",
        "cost",
    ],
    "risk_and_security": [
        "rủi ro",
        "bảo mật",
        "an toàn thông tin",
        "phân quyền",
        "risk",
        "security",
        "privacy",
    ],
    "acceptance_and_testing": [
        "nghiệm thu",
        "kiểm thử",
        "tiêu chí chấp nhận",
        "acceptance criteria",
        "testing",
        "test case",
    ],
    "operation_and_maintenance": [
        "vận hành",
        "bảo trì",
        "bảo hành",
        "hỗ trợ kỹ thuật",
        "maintenance",
        "operation",
    ],
}

NUMBER_PATTERN = re.compile(
    r"""
    (?<!\w)
    [-+]?
    (?:
        \d{1,3}(?:[.,]\d{3})+
        |
        \d+(?:[.,]\d+)?
    )
    \s*
    (?:
        %|％|
        triệu|tỷ|nghìn|ngàn|
        đồng|vnđ|vnd|
        ngày|tháng|năm|tuần|giờ|phút|
        người|đơn vị|hệ thống|máy chủ|
        gb|mb|tb|kb|
        kg|km|m2|m²|
        usd|eur|jpy
    )?
    (?!\w)
    """,
    re.IGNORECASE | re.VERBOSE,
)

DATE_PATTERNS = [
    re.compile(r"\b\d{1,2}[/-]\d{1,2}[/-]\d{2,4}\b"),
    re.compile(r"\b\d{4}[/-]\d{1,2}[/-]\d{1,2}\b"),
    re.compile(
        r"\b(?:ngày|tháng|năm)\s+\d{1,4}\b",
        re.IGNORECASE,
    ),
]

MAX_TEXT_CHARS_PER_BLOCK = 12000


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()

    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)

    return digest.hexdigest()


def normalize_text(text: str) -> str:
    text = text.replace("\xa0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def add_block(
    blocks: list[dict[str, Any]],
    kind: str,
    location: str,
    text: str,
) -> None:
    cleaned = normalize_text(text)

    if not cleaned:
        return

    if len(cleaned) <= MAX_TEXT_CHARS_PER_BLOCK:
        blocks.append(
            {
                "kind": kind,
                "location": location,
                "text": cleaned,
            }
        )
        return

    # Preserve long paragraphs/cells as multiple evidence blocks.
    start = 0
    part = 1

    while start < len(cleaned):
        end = min(start + MAX_TEXT_CHARS_PER_BLOCK, len(cleaned))

        if end < len(cleaned):
            boundary = cleaned.rfind(" ", start, end)

            if boundary > start:
                end = boundary

        blocks.append(
            {
                "kind": kind,
                "location": f"{location} [part {part}]",
                "text": cleaned[start:end].strip(),
            }
        )

        start = end
        part += 1


def extract_document(path: Path) -> dict[str, Any]:
    document = Document(str(path))
    blocks: list[dict[str, Any]] = []

    paragraph_count = 0
    table_cell_count = 0

    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text or ""

        if text.strip():
            paragraph_count += 1

            style_name = ""
            try:
                style_name = paragraph.style.name or ""
            except Exception:
                pass

            location = f"paragraph {index}"

            if style_name:
                location += f" (style: {style_name})"

            add_block(
                blocks=blocks,
                kind="paragraph",
                location=location,
                text=text,
            )

    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            for cell_index, cell in enumerate(row.cells, start=1):
                text = cell.text or ""

                if not text.strip():
                    continue

                table_cell_count += 1

                location = (
                    f"table {table_index}, "
                    f"row {row_index}, "
                    f"cell {cell_index}"
                )

                add_block(
                    blocks=blocks,
                    kind="table_cell",
                    location=location,
                    text=text,
                )

    combined_text = "\n".join(
        block["text"]
        for block in blocks
    )

    return {
        "filename": path.name,
        "path": str(path.resolve()),
        "sha256": sha256_file(path),
        "text": combined_text,
        "word_count": len(combined_text.split()),
        "block_count": len(blocks),
        "paragraph_count": paragraph_count,
        "table_cell_count": table_cell_count,
        "blocks": blocks,
    }


def detect_placeholders(
    document: dict[str, Any],
) -> list[dict[str, Any]]:
    findings: list[dict[str, Any]] = []

    for block in document["blocks"]:
        for pattern in PLACEHOLDER_PATTERNS:
            for match in pattern.finditer(block["text"]):
                findings.append(
                    {
                        "type": "placeholder",
                        "severity": "medium",
                        "filename": document["filename"],
                        "location": block["location"],
                        "quote": match.group(0),
                        "message": (
                            "Phát hiện nội dung có thể chưa hoàn thiện."
                        ),
                        "source": "deterministic",
                    }
                )

    return findings


def detect_missing_content(
    document: dict[str, Any],
) -> list[dict[str, Any]]:
    findings: list[dict[str, Any]] = []
    text = document["text"].casefold()

    for topic, keywords in TOPIC_MAP.items():
        if any(keyword.casefold() in text for keyword in keywords):
            continue

        findings.append(
            {
                "type": "topic_not_detected",
                "severity": "low",
                "filename": document["filename"],
                "location": "whole document",
                "quote": "",
                "topic": topic,
                "message": (
                    f"Không phát hiện từ khóa đại diện cho chủ đề "
                    f"'{topic}'. Đây là tín hiệu để rà soát, "
                    f"không phải bằng chứng chắc chắn rằng tài liệu thiếu."
                ),
                "source": "deterministic",
            }
        )

    return findings


def extract_numeric_mentions(
    document: dict[str, Any],
) -> list[dict[str, Any]]:
    mentions: list[dict[str, Any]] = []

    for block in document["blocks"]:
        text = block["text"]

        for pattern in DATE_PATTERNS:
            for match in pattern.finditer(text):
                mentions.append(
                    {
                        "kind": "date",
                        "value": match.group(0),
                        "filename": document["filename"],
                        "location": block["location"],
                        "context": text[
                            max(0, match.start() - 100):
                            min(len(text), match.end() + 100)
                        ],
                    }
                )

        for match in NUMBER_PATTERN.finditer(text):
            mentions.append(
                {
                    "kind": "number_or_measure",
                    "value": match.group(0).strip(),
                    "filename": document["filename"],
                    "location": block["location"],
                    "context": text[
                        max(0, match.start() - 100):
                        min(len(text), match.end() + 100)
                    ],
                }
            )

    return mentions


def assess_document(
    document: dict[str, Any],
) -> dict[str, Any]:
    findings: list[dict[str, Any]] = []

    findings.extend(detect_placeholders(document))
    findings.extend(detect_missing_content(document))

    topic_coverage: dict[str, Any] = {}
    text = document["text"].casefold()

    for topic, keywords in TOPIC_MAP.items():
        matched = [
            keyword
            for keyword in keywords
            if keyword.casefold() in text
        ]

        topic_coverage[topic] = {
            "detected": bool(matched),
            "matched_keywords": matched,
        }

    return {
        "filename": document["filename"],
        "path": document["path"],
        "sha256": document["sha256"],
        "word_count": document["word_count"],
        "block_count": document["block_count"],
        "paragraph_count": document["paragraph_count"],
        "table_cell_count": document["table_cell_count"],
        "topic_coverage": topic_coverage,
        "findings": findings,
        "numeric_mentions": extract_numeric_mentions(document),
        # Critical: semantic_review.py consumes these evidence blocks.
        "blocks": [
            {
                "kind": block["kind"],
                "location": block["location"],
                "text": block["text"],
            }
            for block in document["blocks"]
        ],
    }


def find_duplicate_groups(
    documents: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    by_hash: dict[str, list[str]] = defaultdict(list)

    for document in documents:
        by_hash[document["sha256"]].append(document["filename"])

    return [
        {
            "filenames": filenames,
            "count": len(filenames),
            "type": "identical_file_content",
        }
        for filenames in by_hash.values()
        if len(filenames) > 1
    ]


def find_cross_document_candidates(
    documents: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """
    Tạo các cặp tài liệu cần đối chiếu sâu hơn.

    Đây chỉ là bước sàng lọc bằng dấu hiệu từ khóa và số liệu.
    Nó chưa khẳng định hai tài liệu mâu thuẫn.
    """
    candidates: list[dict[str, Any]] = []

    for left_index in range(len(documents)):
        left = documents[left_index]

        for right_index in range(left_index + 1, len(documents)):
            right = documents[right_index]

            left_topics = {
                topic
                for topic, info in left["topic_coverage"].items()
                if info["detected"]
            }

            right_topics = {
                topic
                for topic, info in right["topic_coverage"].items()
                if info["detected"]
            }

            shared_topics = sorted(left_topics & right_topics)

            left_numbers = {
                mention["value"].casefold()
                for mention in left["numeric_mentions"]
            }

            right_numbers = {
                mention["value"].casefold()
                for mention in right["numeric_mentions"]
            }

            shared_numbers = sorted(left_numbers & right_numbers)

            candidates.append(
                {
                    "left": left["filename"],
                    "right": right["filename"],
                    "shared_topics": shared_topics,
                    "shared_number_count": len(shared_numbers),
                    "shared_numbers_sample": shared_numbers[:20],
                    "reason": (
                        "shared_topics"
                        if shared_topics
                        else "pairwise_screening"
                    ),
                }
            )

    return candidates


def build_markdown_report(report: dict[str, Any]) -> str:
    lines: list[str] = [
        "# JEV Document Analysis",
        "",
        f"- Generated at: {report['generated_at']}",
        f"- Documents found: {report['documents_found']}",
        f"- Documents selected: {report['documents_selected']}",
        f"- Documents read: {report['documents_read']}",
        f"- Read errors: {report['read_errors']}",
        f"- Findings: {report['finding_count']}",
        f"- Cross-document candidates: "
        f"{len(report['cross_document_candidates'])}",
        "",
        "## Documents",
        "",
    ]

    for document in report["documents"]:
        lines.extend(
            [
                f"### {document['filename']}",
                "",
                f"- Words: {document['word_count']}",
                f"- Evidence blocks: {document['block_count']}",
                f"- Paragraphs: {document['paragraph_count']}",
                f"- Table cells: {document['table_cell_count']}",
                f"- SHA-256: `{document['sha256']}`",
                "",
                "#### Topic coverage",
                "",
            ]
        )

        for topic, info in document["topic_coverage"].items():
            status = "detected" if info["detected"] else "not detected"
            matched = ", ".join(info["matched_keywords"]) or "—"
            lines.append(f"- **{topic}**: {status}; keywords: {matched}")

        lines.extend(["", "#### Findings", ""])

        if not document["findings"]:
            lines.append("- No deterministic findings.")
        else:
            for finding in document["findings"]:
                lines.append(
                    f"- [{finding['severity']}] "
                    f"{finding['type']} — {finding['location']}: "
                    f"{finding['message']}"
                )

        lines.extend(["", "#### Numeric/date mentions", ""])

        if not document["numeric_mentions"]:
            lines.append("- No numeric/date mentions detected.")
        else:
            for mention in document["numeric_mentions"][:100]:
                lines.append(
                    f"- `{mention['value']}` — {mention['location']}: "
                    f"{mention['context']}"
                )

        lines.append("")

    lines.extend(["## Cross-document candidates", ""])

    if not report["cross_document_candidates"]:
        lines.append("- No pairs to compare.")
    else:
        for candidate in report["cross_document_candidates"]:
            topics = ", ".join(candidate["shared_topics"]) or "none detected"
            lines.append(
                f"- **{candidate['left']}** ↔ **{candidate['right']}**; "
                f"shared topics: {topics}; "
                f"shared number strings: "
                f"{candidate['shared_number_count']}"
            )

    lines.extend(
        [
            "",
            "## Limitations",
            "",
            "- Keyword coverage is a screening signal, not proof of completeness.",
            "- Numeric mentions are extracted as text; equivalent values may differ in formatting.",
            "- Cross-document candidates are not confirmed contradictions.",
            "- DOCX text extraction may not include text inside images or scanned pages.",
            "- Header/footer content, text boxes, tracked changes, and embedded objects "
            "may not be fully captured.",
            "",
        ]
    )

    return "\n".join(lines)


def collect_docx_files(
    directory: Path,
    max_documents: int,
) -> tuple[list[Path], int]:
    if not directory.exists():
        return [], 0

    all_files = sorted(
        (
            path
            for path in directory.rglob("*.docx")
            if path.is_file() and not path.name.startswith("~$")
        ),
        key=lambda path: str(path).casefold(),
    )

    return all_files[:max_documents], len(all_files)


def run_analysis(
    documents_dir: Path = DOCUMENTS_DIR,
    reports_dir: Path = REPORTS_DIR,
    max_documents: int = DEFAULT_MAX_DOCUMENTS,
) -> dict[str, Any]:
    if max_documents < 1:
        raise ValueError("max_documents must be at least 1")

    reports_dir.mkdir(parents=True, exist_ok=True)

    selected_paths, documents_found = collect_docx_files(
        documents_dir,
        max_documents,
    )

    assessed_documents: list[dict[str, Any]] = []
    read_errors: list[dict[str, str]] = []

    for path in selected_paths:
        try:
            extracted = extract_document(path)
            assessed_documents.append(assess_document(extracted))
        except Exception as error:
            read_errors.append(
                {
                    "filename": path.name,
                    "path": str(path.resolve()),
                    "error": f"{type(error).__name__}: {error}",
                }
            )

    all_findings = [
        finding
        for document in assessed_documents
        for finding in document["findings"]
    ]

    report: dict[str, Any] = {
        "schema_version": "1.1",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "documents_directory": str(documents_dir.resolve()),
        "reports_directory": str(reports_dir.resolve()),
        "max_documents": max_documents,
        "documents_found": documents_found,
        "documents_selected": len(selected_paths),
        "documents_read": len(assessed_documents),
        "read_errors": len(read_errors),
        "read_error_details": read_errors,
        "finding_count": len(all_findings),
        "documents": assessed_documents,
        "duplicate_groups": find_duplicate_groups(assessed_documents),
        "cross_document_candidates": find_cross_document_candidates(
            assessed_documents
        ),
    }

    json_path = reports_dir / "document_analysis.json"
    markdown_path = reports_dir / "document_analysis.md"

    json_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    markdown_path.write_text(
        build_markdown_report(report),
        encoding="utf-8",
    )

    print("=== JEV DOCUMENT ANALYZER ===")
    print(
        json.dumps(
            {
                "documents_found": report["documents_found"],
                "documents_selected": report["documents_selected"],
                "documents_read": report["documents_read"],
                "read_errors": report["read_errors"],
                "document_findings": report["finding_count"],
                "cross_document_candidates": len(
                    report["cross_document_candidates"]
                ),
                "duplicate_groups": len(report["duplicate_groups"]),
                "evidence_blocks": sum(
                    len(document.get("blocks", []))
                    for document in report["documents"]
                ),
            },
            ensure_ascii=False,
            indent=2,
        )
    )

    print(f"JSON: {json_path}")
    print(f"Markdown: {markdown_path}")

    return report


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Analyze and screen up to 50 DOCX documents for JEV."
    )
    parser.add_argument(
        "--documents-dir",
        type=Path,
        default=DOCUMENTS_DIR,
        help="Directory containing DOCX documents.",
    )
    parser.add_argument(
        "--reports-dir",
        type=Path,
        default=REPORTS_DIR,
        help="Directory for generated reports.",
    )
    parser.add_argument(
        "--max-documents",
        type=int,
        default=DEFAULT_MAX_DOCUMENTS,
        help="Maximum number of DOCX files to process (default: 50).",
    )

    args = parser.parse_args()

    run_analysis(
        documents_dir=args.documents_dir,
        reports_dir=args.reports_dir,
        max_documents=args.max_documents,
    )


if __name__ == "__main__":
    main()