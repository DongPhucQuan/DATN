from __future__ import annotations

import argparse
import json
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook


DEFAULT_FILES = {
    "MRD": Path("BM.01.01.MRD.NV_2023.docx"),
    "PRD": Path("BM.01.02.PRD.NV_2023.docx"),
    "TTR": Path("BM.01.03.TTR.NV_2023.docx"),
    "BUDGET": Path("BM.01_Dự toán.xlsx"),
}

# Các giá trị cần xác minh trực tiếp trên tài liệu.
CHECKS = {
    "duration_months": {
        "label": "Thời gian thực hiện 18 tháng",
        "patterns": [
            r"\b18\s*tháng\b",
            r"\b18\s*months?\b",
        ],
        "documents": ["MRD", "PRD", "TTR"],
    },
    "frequency_n78": {
        "label": "Băng tần n78 / 3.5 GHz",
        "patterns": [
            r"\bn78\b",
            r"3[.,]5\s*GHz",
            r"3500\s*MHz",
        ],
        "documents": ["MRD", "PRD", "TTR"],
    },
    "bandwidth_100mhz": {
        "label": "Băng thông 100 MHz",
        "patterns": [
            r"\b100\s*MHz\b",
        ],
        "documents": ["MRD", "PRD", "TTR"],
    },
    "throughput_1_2gbps": {
        "label": "Thông lượng tối thiểu 1.2 Gbps",
        "patterns": [
            r"1[.,]2\s*Gbps",
            r"1200\s*Mbps",
        ],
        "documents": ["MRD", "PRD", "TTR"],
    },
    "evm_3_2": {
        "label": "EVM không vượt quá 3.2%",
        "patterns": [
            r"3[.,]2\s*%",
            r"3[.,]2\s*percent",
        ],
        "documents": ["MRD", "PRD", "TTR"],
    },
    "runtime_72h": {
        "label": "Chạy liên tục tối thiểu 72 giờ",
        "patterns": [
            r"\b72\s*(?:giờ|h|hours?)\b",
        ],
        "documents": ["MRD", "PRD", "TTR"],
    },
    "rf_4t4r": {
        "label": "Cấu hình nguyên mẫu 4T4R",
        "patterns": [
            r"\b4T4R\b",
        ],
        "documents": ["MRD", "PRD", "TTR"],
    },
    "rf_32t32r": {
        "label": "Cấu hình tham chiếu 32T32R",
        "patterns": [
            r"\b32T32R\b",
        ],
        "documents": ["MRD", "PRD", "TTR"],
    },
    "cost_reduction_40": {
        "label": "Mục tiêu giảm chi phí 40%",
        "patterns": [
            r"\b40\s*%",
            r"\b40\s*percent\b",
        ],
        "documents": ["MRD", "PRD", "TTR"],
    },
}


def normalize_text(value: Any) -> str:
    text = "" if value is None else str(value)
    text = text.replace("\u00a0", " ")
    text = text.replace("–", "-").replace("—", "-")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def read_docx(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        raise FileNotFoundError(f"Không tìm thấy: {path}")

    document = Document(path)
    blocks: list[dict[str, str]] = []

    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = normalize_text(paragraph.text)
        if text:
            blocks.append({
                "document": path.name,
                "location": f"paragraph {index}",
                "text": text,
            })

    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            cells = [
                normalize_text(cell.text)
                for cell in row.cells
            ]
            text = " | ".join(cell for cell in cells if cell)

            if text:
                blocks.append({
                    "document": path.name,
                    "location": (
                        f"table {table_index}, row {row_index}"
                    ),
                    "text": text,
                })

    return blocks


def read_xlsx(path: Path) -> tuple[list[dict[str, str]], dict[str, Any]]:
    if not path.exists():
        raise FileNotFoundError(f"Không tìm thấy: {path}")

    # Đọc giá trị đã lưu trong workbook.
    value_wb = load_workbook(
        path,
        data_only=True,
        read_only=False,
    )

    # Đọc công thức để kiểm tra riêng.
    formula_wb = load_workbook(
        path,
        data_only=False,
        read_only=False,
    )

    blocks: list[dict[str, str]] = []
    formulas: list[dict[str, Any]] = []
    sheets: list[str] = []

    for sheet_name in value_wb.sheetnames:
        sheets.append(sheet_name)
        value_ws = value_wb[sheet_name]
        formula_ws = formula_wb[sheet_name]

        for row in value_ws.iter_rows():
            parts: list[str] = []

            for cell in row:
                if cell.value is None:
                    continue

                address = cell.coordinate
                value = normalize_text(cell.value)
                parts.append(f"{address}={value}")

                formula_cell = formula_ws[address]

                if (
                    isinstance(formula_cell.value, str)
                    and formula_cell.value.startswith("=")
                ):
                    formulas.append({
                        "sheet": sheet_name,
                        "cell": address,
                        "formula": formula_cell.value,
                        "cached_value": cell.value,
                    })

            if parts:
                blocks.append({
                    "document": path.name,
                    "location": f"sheet {sheet_name}",
                    "text": " | ".join(parts),
                })

    value_wb.close()
    formula_wb.close()

    return blocks, {
        "sheets": sheets,
        "formulas": formulas,
    }


def contains_any(text: str, patterns: list[str]) -> bool:
    return any(
        re.search(pattern, text, flags=re.IGNORECASE)
        for pattern in patterns
    )


def run_parameter_checks(
    documents: dict[str, list[dict[str, str]]],
) -> list[dict[str, Any]]:
    results: list[dict[str, Any]] = []

    for check_id, config in CHECKS.items():
        document_results = []

        for document_key in config["documents"]:
            blocks = documents.get(document_key, [])

            matches = [
                {
                    "location": block["location"],
                    "quote": block["text"],
                }
                for block in blocks
                if contains_any(block["text"], config["patterns"])
            ]

            document_results.append({
                "document": document_key,
                "status": "found" if matches else "not_found",
                "matches": matches[:10],
                "match_count": len(matches),
            })

        found_count = sum(
            item["status"] == "found"
            for item in document_results
        )

        if found_count == len(document_results):
            status = "present_in_all"
        elif found_count == 0:
            status = "not_found_in_any"
        else:
            status = "incomplete_coverage"

        results.append({
            "check_id": check_id,
            "label": config["label"],
            "status": status,
            "documents": document_results,
            "interpretation": (
                "Chỉ xác nhận có xuất hiện chuỗi tương ứng; "
                "chưa xác nhận ngữ cảnh, điều kiện đo hoặc ý nghĩa."
            ),
        })

    return results


def parse_decimal(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None

    if isinstance(value, (int, float, Decimal)):
        try:
            return Decimal(str(value))
        except InvalidOperation:
            return None

    text = normalize_text(value)
    if not text:
        return None

    # Chỉ nhận ô chứa một số đơn giản.
    # Không tự đoán định dạng số có dấu chấm/phẩy hàng nghìn.
    if re.fullmatch(r"[-+]?\d+(?:[.,]\d+)?", text):
        try:
            return Decimal(text.replace(",", "."))
        except InvalidOperation:
            return None

    return None


def audit_workbook_formulas(path: Path) -> dict[str, Any]:
    """
    Kiểm tra công thức có tồn tại và có giá trị cache hay không.
    Không tự tính lại công thức Excel.
    """
    formula_wb = load_workbook(
        path,
        data_only=False,
        read_only=False,
    )
    value_wb = load_workbook(
        path,
        data_only=True,
        read_only=False,
    )

    formula_count = 0
    missing_cached_values = []
    formula_details = []

    for sheet_name in formula_wb.sheetnames:
        formula_ws = formula_wb[sheet_name]
        value_ws = value_wb[sheet_name]

        for row in formula_ws.iter_rows():
            for cell in row:
                formula = cell.value

                if not (
                    isinstance(formula, str)
                    and formula.startswith("=")
                ):
                    continue

                formula_count += 1
                cached_value = value_ws[cell.coordinate].value

                detail = {
                    "sheet": sheet_name,
                    "cell": cell.coordinate,
                    "formula": formula,
                    "cached_value": cached_value,
                }
                formula_details.append(detail)

                if cached_value is None:
                    missing_cached_values.append({
                        "sheet": sheet_name,
                        "cell": cell.coordinate,
                        "formula": formula,
                    })

    formula_wb.close()
    value_wb.close()

    return {
        "formula_count": formula_count,
        "missing_cached_value_count": len(missing_cached_values),
        "missing_cached_values": missing_cached_values,
        "formulas": formula_details,
        "limitation": (
            "openpyxl không tính lại công thức Excel. "
            "Giá trị cache có thể cũ hoặc chưa được lưu."
        ),
    }


def inspect_budget_total(
    blocks: list[dict[str, str]],
) -> dict[str, Any]:
    """
    Tìm các dòng có từ khóa tổng kinh phí/tổng cộng.
    Không cộng tùy tiện tất cả ô số vì có thể đếm trùng tổng phụ.
    """
    keywords = [
        "tổng kinh phí",
        "tổng cộng",
        "kinh phí tổng",
        "tổng dự toán",
    ]

    matches = []

    for block in blocks:
        text = block["text"]
        normalized = text.casefold()

        if any(keyword in normalized for keyword in keywords):
            matches.append({
                "location": block["location"],
                "text": text,
            })

    return {
        "status": "candidates_found" if matches else "not_found",
        "candidate_count": len(matches),
        "candidates": matches[:100],
        "note": (
            "Phải xác định đúng dòng tổng và đơn vị tính trước khi "
            "kết luận tổng dự toán. Không cộng tất cả số tìm thấy."
        ),
    }


def audit_document_presence(
    documents: dict[str, list[dict[str, str]]],
) -> dict[str, Any]:
    return {
        key: {
            "block_count": len(blocks),
            "empty": len(blocks) == 0,
        }
        for key, blocks in documents.items()
    }


def build_report(
    files: dict[str, Path],
) -> dict[str, Any]:
    documents: dict[str, list[dict[str, str]]] = {}

    for key in ("MRD", "PRD", "TTR"):
        documents[key] = read_docx(files[key])

    budget_blocks, workbook_meta = read_xlsx(files["BUDGET"])
    documents["BUDGET"] = budget_blocks

    parameter_results = run_parameter_checks(documents)

    report = {
        "engine": "JEV Deterministic Cross-Document Audit",
        "mode": "read-only; no output files",
        "inputs": {
            key: str(path.resolve())
            for key, path in files.items()
        },
        "document_summary": audit_document_presence(documents),
        "parameter_checks": parameter_results,
        "budget_total_candidates": inspect_budget_total(
            budget_blocks
        ),
        "workbook": {
            "sheets": workbook_meta["sheets"],
            "formula_audit": audit_workbook_formulas(
                files["BUDGET"]
            ),
        },
        "limitations": [
            "Chuỗi xuất hiện không chứng minh hai yêu cầu tương đương.",
            "Không tìm thấy chuỗi không chứng minh nội dung bị thiếu.",
            "Không tự kết luận mâu thuẫn nếu chưa xét ngữ cảnh.",
            "Chưa tự tính lại công thức Excel.",
            "Tổng dự toán phải được xác minh theo dòng tổng và đơn vị.",
            "Mọi kết luận ngữ nghĩa cần người kiểm tra.",
        ],
    }

    return report


def main() -> None:
    parser = argparse.ArgumentParser(
        description="JEV read-only audit for 3 DOCX + 1 XLSX."
    )

    parser.add_argument(
        "--mrd",
        type=Path,
        default=DEFAULT_FILES["MRD"],
    )
    parser.add_argument(
        "--prd",
        type=Path,
        default=DEFAULT_FILES["PRD"],
    )
    parser.add_argument(
        "--ttr",
        type=Path,
        default=DEFAULT_FILES["TTR"],
    )
    parser.add_argument(
        "--budget",
        type=Path,
        default=DEFAULT_FILES["BUDGET"],
    )

    args = parser.parse_args()

    files = {
        "MRD": args.mrd,
        "PRD": args.prd,
        "TTR": args.ttr,
        "BUDGET": args.budget,
    }

    try:
        report = build_report(files)
    except Exception as exc:
        raise SystemExit(
            f"[JEV AUDIT ERROR] {type(exc).__name__}: {exc}"
        ) from exc

    print(json.dumps(
        report,
        ensure_ascii=False,
        indent=2,
        default=str,
    ))


if __name__ == "__main__":
    main()