
from __future__ import annotations

import hashlib
import json
import sys
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
INPUT_DIR = ROOT / "data" / "documents"
OUTPUT_FILE = ROOT / "reports" / "document_analysis.json"

EXPECTED_DOCUMENTS = 4
SUPPORTED_EXTENSIONS = {".docx", ".xlsx"}
CHUNK_SIZE = 1800


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def split_text(text: str, limit: int = CHUNK_SIZE) -> list[str]:
    text = " ".join(text.split())
    if not text:
        return []

    chunks = []
    while len(text) > limit:
        split_at = text.rfind(" ", 0, limit + 1)
        if split_at < limit // 2:
            split_at = limit
        chunks.append(text[:split_at].strip())
        text = text[split_at:].strip()

    if text:
        chunks.append(text)

    return chunks


def make_blocks(
    text: str,
    file_name: str,
    file_type: str,
    location: str,
    block_prefix: str,
) -> list[dict[str, Any]]:
    blocks = []

    for index, chunk in enumerate(split_text(text), start=1):
        blocks.append({
            "block_id": f"{block_prefix}_{index:04d}",
            "source_file": file_name,
            "source_type": file_type,
            "location": location,
            "text": chunk,
        })

    return blocks


def read_docx(path: Path) -> list[dict[str, Any]]:
    document = Document(path)
    blocks = []

    for paragraph_index, paragraph in enumerate(
        document.paragraphs, start=1
    ):
        text = paragraph.text.strip()
        if text:
            blocks.extend(
                make_blocks(
                    text=text,
                    file_name=path.name,
                    file_type="docx",
                    location=f"paragraph:{paragraph_index}",
                    block_prefix=f"{path.stem}_P{paragraph_index}",
                )
            )

    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            cells = [
                " ".join(cell.text.split())
                for cell in row.cells
            ]
            if not any(cells):
                continue

            blocks.extend(
                make_blocks(
                    text=" | ".join(cells),
                    file_name=path.name,
                    file_type="docx",
                    location=f"table:{table_index}/row:{row_index}",
                    block_prefix=(
                        f"{path.stem}_T{table_index}_R{row_index}"
                    ),
                )
            )

    return blocks


def read_xlsx(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(
        path,
        read_only=True,
        data_only=True,
    )
    blocks = []

    try:
        for sheet in workbook.worksheets:
            for row_index, row in enumerate(
                sheet.iter_rows(values_only=True), start=1
            ):
                values = [normalize(value) for value in row]

                if not any(
                    value is not None and str(value).strip()
                    for value in values
                ):
                    continue

                cells = [
                    "" if value is None else str(value)
                    for value in values
                ]

                blocks.extend(
                    make_blocks(
                        text=" | ".join(cells),
                        file_name=path.name,
                        file_type="xlsx",
                        location=(
                            f"sheet:{sheet.title}/row:{row_index}"
                        ),
                        block_prefix=(
                            f"{path.stem}_{sheet.title}_R{row_index}"
                        ),
                    )
                )
    finally:
        workbook.close()

    return blocks


def collect_documents() -> list[Path]:
    if not INPUT_DIR.is_dir():
        raise FileNotFoundError(
            f"Input directory does not exist: {INPUT_DIR.resolve()}"
        )

    return sorted(
        path for path in INPUT_DIR.iterdir()
        if path.is_file()
        and path.suffix.lower() in SUPPORTED_EXTENSIONS
        and not path.name.startswith("~$")
    )


def main() -> int:
    files = collect_documents()

    print("=== JEV DOCUMENT INPUT AUDIT ===")
    print(f"Project: {ROOT.resolve()}")
    print(f"Input:   {INPUT_DIR.resolve()}")
    print(f"Output:  {OUTPUT_FILE.resolve()}")
    print(f"Expected document count: {EXPECTED_DOCUMENTS}")
    print(f"Discovered document count: {len(files)}")

    if len(files) != EXPECTED_DOCUMENTS:
        print("[ERROR] Input file count mismatch; no output written.")
        for path in files:
            print(f"  Found: {path.name}")
        return 1

    documents = []
    all_blocks = []

    for path in files:
        file_info = {
            "file_name": path.name,
            "absolute_path": str(path.resolve()),
            "file_type": path.suffix.lower(),
            "sha256": sha256_file(path),
            "size_bytes": path.stat().st_size,
            "modified_time_utc": datetime.fromtimestamp(
                path.stat().st_mtime,
                tz=timezone.utc,
            ).isoformat(),
        }

        if path.suffix.lower() == ".docx":
            blocks = read_docx(path)
        elif path.suffix.lower() == ".xlsx":
            blocks = read_xlsx(path)
        else:
            continue

        if not blocks:
            print(f"[ERROR] No readable text found: {path.name}")
            return 1

        file_info["block_count"] = len(blocks)
        documents.append(file_info)
        all_blocks.extend(blocks)

        print(f"\n[READ] {path.name}")
        print(f"  Path:   {path.resolve()}")
        print(f"  SHA256: {file_info['sha256']}")
        print(f"  Bytes:  {file_info['size_bytes']}")
        print(f"  Blocks: {len(blocks)}")

    docx_count = sum(
        item["file_type"] == ".docx" for item in documents
    )
    xlsx_count = sum(
        item["file_type"] == ".xlsx" for item in documents
    )

    if docx_count != 3 or xlsx_count != 1:
        print(
            "\n[ERROR] Expected exactly 3 DOCX and 1 XLSX; "
            "no output written."
        )
        print(f"Found DOCX: {docx_count}; XLSX: {xlsx_count}")
        return 1

    result = {
        "schema_version": 1,
        "input_mode": "real",
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "project_root": str(ROOT.resolve()),
        "source_directory": str(INPUT_DIR.resolve()),
        "document_count": len(documents),
        "block_count": len(all_blocks),
        "source_files": documents,
        "documents": [
            {
                "document_id": Path(item["file_name"]).stem,
                "file_name": item["file_name"],
                "file_type": item["file_type"],
                "block_count": item["block_count"],
            }
            for item in documents
        ],
        "evidence_blocks": all_blocks,
    }

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    temporary = OUTPUT_FILE.with_suffix(".json.tmp")

    try:
        temporary.write_text(
            json.dumps(result, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        temporary.replace(OUTPUT_FILE)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise

    print("\n=== INGESTION COMPLETE ===")
    print(f"Documents: {len(documents)}")
    print(f"Evidence blocks: {len(all_blocks)}")
    print(f"Output: {OUTPUT_FILE.resolve()}")
    print("Original source files were not modified.")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:
        print(f"[ERROR] {type(exc).__name__}: {exc}", file=sys.stderr)
        sys.exit(1)
